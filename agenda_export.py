"""
agenda_export.py - Turn agenda days into files a substitute or OneNote can use.

Two renderers over the same page payloads (built by ui/agendas_view.py):

  * write_pdf(pages, path)   — looks like the projected day: the date up top,
                               reminders/announcements, the percussion (or jazz
                               rhythm) rotation when one is filled in, then the
                               sections with an empty check box per line.
  * write_docx(pages, path)  — the same day as an editable Word document, so a
                               teacher can add notes for the sub before printing.

One page per (day, class period).  No tkinter here — pure data in, file out,
so the whole thing is testable headless.

A page payload:
    {label, date (datetime.date), period ("Period 6" or ""),
     reminders: [str], announcements: [str],
     perc_head: str, perc_rows: [(name, station)],
     sections: [{title, items: [
         {text, runs, indent, static, color}      — a line
         {image: abspath, img_w: px}              — a picture
     ]}]}

``runs`` are the agenda's rich-text spans [[start, end, tag]] with tag in
b / i / u / hl — the same shape the editor stores.
"""

from __future__ import annotations

import os

# Line colors, matching ui/agendas_view.py.  "white" is projector-only (white
# text on a dark screen); on paper it prints as plain ink.
_TEXT_HEX = {"black": "#111111", "blue": "#1565d8", "red": "#e11414"}
_HL_BG = "#fff59d"
_HDR_BG = "#3b7dc4"
_BAN_BG = "#FFF3BF"
_BAN_FG = "#5C4A00"


def _fmt_date(d):
    return d.strftime("%A, %B ") + f"{d.day}, {d.year}"


def _segment(text, runs):
    """Split a line into (chunk, {tags}) pieces so overlapping bold/italic/
    underline/highlight spans come out right in both renderers."""
    text = text or ""
    marks = [set() for _ in range(len(text))]
    for r in (runs or []):
        try:
            s, e, tag = int(r[0]), int(r[1]), str(r[2])
        except (ValueError, TypeError, IndexError):
            continue
        for i in range(max(0, s), min(len(text), e)):
            marks[i].add(tag)
    out = []
    for i, ch in enumerate(text):
        if out and out[-1][1] == marks[i]:
            out[-1][0] += ch
        else:
            out.append([ch, marks[i]])
    return [(chunk, tags) for chunk, tags in out] or [(text, set())]


# ══════════════════════════════════════════════════════════════════ PDF ═════

_SYM_FONT = None


def _symbol_font():
    """A font that really has ☐ and ◦ — Helvetica doesn't.  Segoe UI Symbol
    ships with every Windows Roka runs on; anything missing falls back to
    plain characters."""
    global _SYM_FONT
    if _SYM_FONT is not None:
        return _SYM_FONT
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        path = os.path.join(os.environ.get("WINDIR", r"C:\Windows"),
                            "Fonts", "seguisym.ttf")
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont("RokaSymbols", path))
            _SYM_FONT = "RokaSymbols"
        else:
            _SYM_FONT = ""
    except Exception:
        _SYM_FONT = ""
    return _SYM_FONT


def _esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;"))


def _para_markup(item):
    """Reportlab paragraph XML for one agenda line, honoring runs + color."""
    color = (item.get("color") or "").strip()
    parts = []
    for chunk, tags in _segment(item.get("text") or "", item.get("runs")):
        s = _esc(chunk)
        if "hl" in tags:
            s = f'<font backColor="{_HL_BG}">{s}</font>'
        if "u" in tags:
            s = f"<u>{s}</u>"
        if "i" in tags:
            s = f"<i>{s}</i>"
        if "b" in tags:
            s = f"<b>{s}</b>"
        parts.append(s)
    body = "".join(parts)
    if color == "hl":
        body = f'<font backColor="{_HL_BG}">{body}</font>'
    elif color in _TEXT_HEX:
        body = f'<font color="{_TEXT_HEX[color]}">{body}</font>'
    return body


def write_pdf(pages, path):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Table,
                                    TableStyle, Spacer, PageBreak,
                                    Image as RLImage)

    sym = _symbol_font()
    box = (f'<font name="{sym}">☐</font>&nbsp;&nbsp;' if sym
           else "[&nbsp;&nbsp;]&nbsp;")
    dot = (f'<font name="{sym}">◦</font>&nbsp;&nbsp;' if sym else "-&nbsp;")

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=11,
                          leading=15.5)
    detail = ParagraphStyle("detail", parent=body, fontSize=10, leading=14,
                            leftIndent=26, textColor=colors.HexColor("#333333"))
    small = ParagraphStyle("small", parent=body, fontSize=9.5, leading=13)
    title = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=16,
                           leading=19)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=11, leading=14,
                         textColor=colors.HexColor("#555555"))
    sechead = ParagraphStyle("sechead", fontName="Helvetica-Bold", fontSize=12,
                             leading=15, textColor=colors.white)
    banhead = ParagraphStyle("banhead", fontName="Helvetica-Bold", fontSize=10,
                             leading=13, textColor=colors.HexColor(_BAN_FG))

    doc = SimpleDocTemplate(path, pagesize=letter,
                            leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.55 * inch, bottomMargin=0.55 * inch)
    avail = doc.width
    story = []

    for pi, page in enumerate(pages):
        if pi:
            story.append(PageBreak())
        head = _esc(page.get("label") or "")
        when = _fmt_date(page["date"])
        story.append(Paragraph(f"{head} — {when}", title))
        if page.get("period"):
            story.append(Paragraph(_esc(page["period"]), sub))
        story.append(Spacer(1, 8))

        # ── banner ──
        rem = page.get("reminders") or []
        ann = page.get("announcements") or []
        if rem or ann:
            def col(head_text, lines):
                out = [Paragraph(head_text, banhead)]
                for ln in lines:
                    out.append(Paragraph("•  " + _esc(ln), small))
                return out
            cells = []
            if rem:
                cells.append(col("Reminders", rem))
            if ann:
                cells.append(col("Announcements", ann))
            t = Table([cells], colWidths=[avail / len(cells)] * len(cells))
            t.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFBEA")),
                ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#E0C048")),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))

        # ── percussion / rhythm rotation ──
        if page.get("perc_rows"):
            data = [[Paragraph(_esc(page.get("perc_head") or "Percussion"),
                               banhead), ""]]
            for name, station in page["perc_rows"]:
                data.append([Paragraph(_esc(name), small),
                             Paragraph(_esc(station), small)])
            t = Table(data, colWidths=[2.0 * inch, avail - 2.0 * inch])
            t.setStyle(TableStyle([
                ("SPAN", (0, 0), (1, 0)),
                ("BACKGROUND", (0, 0), (1, 0), colors.HexColor(_BAN_BG)),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9CFA0")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

        # ── sections ──
        for sec in page.get("sections") or []:
            bar = Table([[Paragraph(_esc(sec.get("title") or ""), sechead)]],
                        colWidths=[avail])
            bar.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(_HDR_BG)),
                ("LEFTPADDING", (0, 0), (0, 0), 8),
                ("TOPPADDING", (0, 0), (0, 0), 4),
                ("BOTTOMPADDING", (0, 0), (0, 0), 4),
            ]))
            story.append(Spacer(1, 6))
            story.append(bar)
            story.append(Spacer(1, 4))
            for item in sec.get("items") or []:
                if item.get("image"):
                    try:
                        reader = ImageReader(item["image"])
                        iw, ih = reader.getSize()
                        w = min(float(item.get("img_w") or 380) * 0.75,
                                avail - 20)
                        story.append(RLImage(item["image"], width=w,
                                             height=ih * (w / iw)))
                        story.append(Spacer(1, 4))
                    except Exception:
                        pass
                    continue
                markup = _para_markup(item)
                if item.get("indent"):
                    story.append(Paragraph(dot + markup, detail))
                elif item.get("static"):
                    story.append(Paragraph(markup, detail))
                else:
                    story.append(Paragraph(box + markup, body))
                story.append(Spacer(1, 2))

    doc.build(story)
    return path


# ═════════════════════════════════════════════════════════════════ DOCX ═════

def _shade(par_or_cell, hex_fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = par_or_cell._p if hasattr(par_or_cell, "_p") else par_or_cell._tc
    pr = (el.get_or_add_pPr() if hasattr(par_or_cell, "_p")
          else el.get_or_add_tcPr())
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill.lstrip("#"))
    pr.append(shd)


def _add_runs(par, item):
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import RGBColor
    color = (item.get("color") or "").strip()
    for chunk, tags in _segment(item.get("text") or "", item.get("runs")):
        run = par.add_run(chunk)
        run.bold = "b" in tags
        run.italic = "i" in tags
        run.underline = "u" in tags
        if "hl" in tags or color == "hl":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if color in _TEXT_HEX:
            run.font.color.rgb = RGBColor.from_string(
                _TEXT_HEX[color].lstrip("#").upper())


def write_docx(pages, path):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for pi, page in enumerate(pages):
        if pi:
            doc.add_page_break()
        p = doc.add_paragraph()
        r = p.add_run(f"{page.get('label') or ''} — {_fmt_date(page['date'])}")
        r.bold = True
        r.font.size = Pt(16)
        if page.get("period"):
            sp = doc.add_paragraph()
            sr = sp.add_run(page["period"])
            sr.font.size = Pt(11)
            sr.font.color.rgb = RGBColor.from_string("555555")

        rem = page.get("reminders") or []
        ann = page.get("announcements") or []
        if rem or ann:
            cols = [c for c in (("Reminders", rem), ("Announcements", ann))
                    if c[1]]
            t = doc.add_table(rows=1, cols=len(cols))
            t.style = "Table Grid"
            for ci, (head, lines) in enumerate(cols):
                cell = t.rows[0].cells[ci]
                _shade(cell, "FFFBEA")
                hp = cell.paragraphs[0]
                hr = hp.add_run(head)
                hr.bold = True
                hr.font.color.rgb = RGBColor.from_string(
                    _BAN_FG.lstrip("#").upper())
                for ln in lines:
                    cell.add_paragraph("•  " + str(ln))

        if page.get("perc_rows"):
            doc.add_paragraph()
            t = doc.add_table(rows=1, cols=2)
            t.style = "Table Grid"
            hcell = t.rows[0].cells[0]
            hcell.merge(t.rows[0].cells[1])
            _shade(hcell, _BAN_BG.lstrip("#"))
            hr = hcell.paragraphs[0].add_run(page.get("perc_head")
                                             or "Percussion")
            hr.bold = True
            for name, station in page["perc_rows"]:
                row = t.add_row()
                row.cells[0].paragraphs[0].add_run(str(name)).bold = True
                row.cells[1].paragraphs[0].add_run(str(station))

        for sec in page.get("sections") or []:
            doc.add_paragraph()
            hp = doc.add_paragraph()
            _shade(hp, _HDR_BG.lstrip("#"))
            hr = hp.add_run(" " + (sec.get("title") or ""))
            hr.bold = True
            hr.font.size = Pt(12)
            hr.font.color.rgb = RGBColor.from_string("FFFFFF")
            for item in sec.get("items") or []:
                if item.get("image"):
                    try:
                        w_in = min(float(item.get("img_w") or 380), 620) / 96.0
                        doc.add_picture(item["image"], width=Inches(w_in))
                    except Exception:
                        pass
                    continue
                par = doc.add_paragraph()
                if item.get("indent"):
                    par.paragraph_format.left_indent = Inches(0.4)
                    par.add_run("◦  ")
                elif item.get("static"):
                    par.paragraph_format.left_indent = Inches(0.25)
                else:
                    par.add_run("☐  ")
                _add_runs(par, item)

    doc.save(path)
    return path
